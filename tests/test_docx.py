from dataclasses import replace

from docx import Document

from kinolist.docx_out import txt_path_for, write_simple_list, write_table_list, write_txt_list
from kinolist.resources import template_path


def test_write_table_list_clones_tables(tmp_path, film):
    output = tmp_path / "list.docx"
    films = [film, replace(film, title="Терминатор 2", rating="", directors=[])]
    assert write_table_list(films, str(output), template_path(a5=False), genres=True)

    document = Document(str(output))
    assert len(document.tables) == 2
    first, second = document.tables
    assert first.cell(0, 1).text == "Терминатор - Кинопоиск 8.0"
    assert "Режиссер: Джеймс Кэмерон" in first.cell(1, 1).text
    assert "Жанр: фантастика" in first.cell(1, 1).text
    assert "В главных ролях: Арнольд Шварценеггер" in first.cell(1, 1).text
    assert second.cell(0, 1).text == "Терминатор 2 - нет рейтинга"
    assert len(document.inline_shapes) == 2


def test_write_table_list_a5(tmp_path, film):
    output = tmp_path / "a5.docx"
    assert write_table_list([film], str(output), template_path(a5=True))
    assert round(Document(str(output)).sections[0].page_height.cm, 1) == 14.8


def test_write_simple_list(tmp_path, film):
    output = tmp_path / "simple.docx"
    assert write_simple_list([film, film], str(output), genres=True)
    paragraphs = [p.text for p in Document(str(output)).paragraphs]
    assert paragraphs[0].startswith("1. Терминатор (1984) Жанр: фантастика\nРежиссер: Джеймс Кэмерон\nАктеры: ")
    assert paragraphs[0].count(",") == 2
    assert paragraphs[1].startswith("2. ")


def test_write_txt_list(tmp_path, film):
    output = str(tmp_path / "list.docx")
    write_txt_list([film, replace(film, title="Другой")], txt_path_for(output))
    assert (tmp_path / "list.txt").read_text(encoding="utf-8") == "Терминатор\nДругой\n"
