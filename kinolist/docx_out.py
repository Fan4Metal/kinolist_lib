"""Формирование списков фильмов: docx по шаблону, docx в простом формате, txt."""

from __future__ import annotations

import logging
import os
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from tqdm import tqdm

from .files import write_lines
from .models import Film, image_to_png
from .resources import no_poster

log = logging.getLogger(__name__)

FONT = "Arial"
TITLE_SIZE = Pt(11)
TEXT_SIZE = Pt(10)
ACTORS_LABEL_COLOR = RGBColor(255, 102, 0)
ACTORS_COLOR = RGBColor(0, 0, 255)
POSTER_WIDTH = Cm(7)
SIMPLE_FONT = "Times New Roman"
SIMPLE_FONT_SIZE = Pt(14)
SIMPLE_ACTORS_LIMIT = 3


def _clone_first_table(document, count: int) -> None:
    """Клонирует первую таблицу документа ``count`` раз, разделяя копии пустыми абзацами."""
    template = document.tables[0]._tbl
    paragraph = document.paragraphs[0]
    for _ in range(count):
        paragraph._p.addnext(deepcopy(template))
        paragraph = document.add_paragraph()


def _add_run(paragraph, text: str, size=TEXT_SIZE, bold=False, color: RGBColor | None = None, underline=False):
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = size
    run.font.bold = bold
    run.font.underline = underline
    if color is not None:
        run.font.color.rgb = color
    return run


def _fill_table(table: Table, film: Film, genres: bool) -> None:
    """Заполняет одну таблицу шаблона данными фильма."""
    _add_run(table.cell(0, 1).paragraphs[0], f"{film.title} - {film.rating_text}", size=TITLE_SIZE, bold=True)

    info = table.cell(1, 1)
    _add_run(info.add_paragraph(), str(film.year) if film.year else "")
    _add_run(info.add_paragraph(), ", ".join(film.countries))
    _add_run(info.add_paragraph(), film.directors_text)
    if genres and film.main_genre:
        _add_run(info.add_paragraph(), f"Жанр: {film.main_genre}")

    info.add_paragraph()
    paragraph = info.add_paragraph()
    _add_run(paragraph, "В главных ролях: ", color=ACTORS_LABEL_COLOR)
    _add_run(paragraph, ", ".join(film.actors), color=ACTORS_COLOR, underline=True)

    info.add_paragraph()
    info.add_paragraph()
    _add_run(info.add_paragraph(), film.description)
    info.add_paragraph()

    poster = film.poster or no_poster()
    table.cell(0, 0).paragraphs[1].add_run().add_picture(image_to_png(poster), width=POSTER_WIDTH)


def _save(document, path: str) -> bool:
    try:
        document.save(path)
    except PermissionError:
        log.error(f'Ошибка! Нет доступа на запись к файлу "{path}". Список не сохранен.')
        return False
    log.info(f'Файл "{path}" создан.')
    return True


def write_table_list(films: list[Film], path: str, template: Path, genres: bool = False) -> bool:
    """Список с постерами по шаблону docx (одна таблица на фильм)."""
    document = Document(str(template))
    if len(films) > 1:
        _clone_first_table(document, len(films) - 1)
    for table, film in zip(tqdm(document.tables, desc="Запись в таблицу...      "), films, strict=True):
        _fill_table(table, film, genres)
    return _save(document, path)


def write_simple_list(films: list[Film], path: str, genres: bool = False) -> bool:
    """Нумерованный текстовый список без постеров («новый формат»)."""
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.75)
    section.bottom_margin = Cm(2)

    font = document.styles["Normal"].font
    font.name = SIMPLE_FONT
    font.size = SIMPLE_FONT_SIZE

    for number, film in enumerate(tqdm(films, desc="Запись в файл...         "), start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(12)
        paragraph.add_run(f"{number}. ")
        paragraph.add_run(f"{film.title} ({film.year}) ").bold = True
        if genres and film.main_genre:
            paragraph.add_run(f"Жанр: {film.main_genre}\n")
        paragraph.add_run(f"{film.directors_text}\n")
        paragraph.add_run(f"Актеры: {', '.join(film.actors[:SIMPLE_ACTORS_LIMIT])}")
    return _save(document, path)


def write_txt_list(films: list[Film], path: str) -> None:
    """Текстовый файл с названиями фильмов, по одному в строке."""
    write_lines(path, [film.title for film in films])
    log.info(f'Файл "{path}" создан.')


def txt_path_for(docx_path: str) -> str:
    return os.path.splitext(docx_path)[0] + ".txt"
